#!/usr/bin/env python3
"""Import filled decision rows from the workbook into walk-<layer>.jsonl.

    python3 sheet-import.py --layer api [--sheet data/triage/reflect/clef/walk-sheet.xlsx] [--by "user"] [--dry-run]

Rows whose Outcome is set to keep | delete | rename | reverse | retarget are written (or
replaced, by triple) in walk-<layer>.jsonl; rows left blank or set to defer are skipped. The
Replacement and Also delete columns must hold 5-part triples ({layer}.{type}.{predicate}.{layer}.{type}).
Then apply with: bun apply-walk.ts --layer <id>
"""
import argparse
import json
import os
import sys
from datetime import date

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
REFL = os.path.join(HERE, "data", "triage", "reflect", "clef")
OUTCOMES = {"keep", "delete", "rename", "reverse", "retarget"}
NEEDS_REPLACEMENT = {"rename", "reverse", "retarget"}


def is_triple(s):
    p = s.split(".")
    return len(p) == 5 and all(p)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--layer", required=True)
    ap.add_argument("--sheet", default=os.path.join(REFL, "walk-sheet.xlsx"))
    ap.add_argument("--by", default="user")
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    wb = load_workbook(args.sheet, data_only=True)
    if args.layer not in wb.sheetnames:
        print(f"no sheet named {args.layer} in {args.sheet}", file=sys.stderr)
        return 2
    ws = wb[args.layer]
    headers = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    col = {h: i for i, h in enumerate(headers)}
    for h in ("Row", "Kind", "Current relationship", "Outcome", "Replacement", "Also delete", "Notes"):
        if h not in col:
            print(f"missing column '{h}'", file=sys.stderr)
            return 2

    out_path = os.path.join(REFL, f"walk-{args.layer}.jsonl")
    existing = {}
    if os.path.exists(out_path):
        with open(out_path, encoding="utf-8") as f:
            for line in f:
                if line.strip():
                    w = json.loads(line)
                    existing[w["triple"]] = w

    errors, new_rows, replaced, skipped = [], 0, 0, 0
    # walk rows are numbered in the order they were decided; sheet rows keep their number if
    # already decided and continue the sequence otherwise
    next_row = max([w["row"] for w in existing.values()], default=0) + 1
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[col["Current relationship"]] is None:
            continue
        triple = str(r[col["Current relationship"]]).strip()
        outcome = (str(r[col["Outcome"]]).strip().lower() if r[col["Outcome"]] is not None else "")
        if outcome not in OUTCOMES:
            skipped += 1
            continue
        replacement = str(r[col["Replacement"]]).strip() if r[col["Replacement"]] is not None else ""
        replacement = replacement.replace(" (exists)", "") or None
        also = [s.strip() for s in str(r[col["Also delete"]] or "").split(",") if s.strip()]
        notes = str(r[col["Notes"]]).strip() if r[col["Notes"]] is not None else ""
        if outcome in NEEDS_REPLACEMENT and not replacement:
            errors.append(f"{triple}: outcome {outcome} needs a Replacement")
            continue
        for t in ([replacement] if replacement else []) + also:
            if not is_triple(t):
                errors.append(f"{triple}: '{t}' is not a 5-part triple")
        if outcome == "keep":
            replacement = None
        if triple in existing:
            row_no = existing[triple]["row"]
        else:
            row_no, next_row = next_row, next_row + 1
        row = {
            "row": row_no,
            "triple": triple,
            "kind": str(r[col["Kind"]]),
            "decision": outcome,
            "replacement": replacement,
            "rationale": notes or f"decided in walk-sheet.xlsx ({outcome})",
            "adjudicated_by": f"{args.by} ({date.today().isoformat()}) via walk-sheet.xlsx",
        }
        if also:
            # extra triples: deletions for delete/keep rows, additions for the rest (apply-walk's `add`)
            row["also_delete" if outcome in ("delete", "keep") else "add"] = also
        if triple in existing:
            replaced += 1
        else:
            new_rows += 1
        existing[triple] = row

    if errors:
        print("not written; fix these cells first:", file=sys.stderr)
        for e in errors:
            print("  " + e, file=sys.stderr)
        return 1
    rows = sorted(existing.values(), key=lambda w: w["row"])
    if not args.dry_run:
        with open(out_path, "w", encoding="utf-8") as f:
            for w in rows:
                f.write(json.dumps(w, ensure_ascii=False) + "\n")
    print(f"{'[dry run] ' if args.dry_run else ''}{args.layer}: {new_rows} new, {replaced} replaced, {skipped} blank or deferred -> {len(rows)} rows in {os.path.relpath(out_path, HERE)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())

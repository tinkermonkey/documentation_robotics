#!/usr/bin/env python3
"""Export the reflection change plan as a decision workbook, one sheet per layer.

    python3 sheet-export.py [--layers api,ux] [--out data/triage/reflect/clef/walk-sheet.xlsx]

Every row of plan-clef.json for the chosen layers (default: every layer that has no
applied-walk-<layer>.json yet) becomes a row with the framing a reviewer needs: the current
relationship and its description, the other edges the catalog has between the two types, the
claims, the three sorting answers, the plan's implied replacement and a mechanical suggestion.
The reviewer fills Outcome (dropdown), Replacement, Also delete and Notes; sheet-import.py
turns the filled rows into walk-<layer>.jsonl for apply-walk.ts.

Rows already decided in walk-<layer>.jsonl are pre-filled from it.
"""
import argparse
import json
import os
import sys
from collections import defaultdict
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(HERE, "..", "..", ".."))
REFL = os.path.join(HERE, "data", "triage", "reflect", "clef")
CLAIMS = os.path.join(HERE, "data", "triage", "claims.jsonl")
REL_ROOT = os.path.join(REPO, "spec", "schemas", "relationships")

OUTCOMES = ["keep", "delete", "rename", "reverse", "retarget", "defer"]
FONT = "Arial"


def read_jsonl(path):
    if not os.path.exists(path):
        return []
    with open(path, encoding="utf-8") as f:
        return [json.loads(line) for line in f if line.strip()]


def parts(triple):
    p = triple.split(".")
    return f"{p[0]}.{p[1]}", p[2], f"{p[3]}.{p[4]}"


def load_catalog():
    """id -> description; plus (source, destination) -> [predicates]"""
    desc, pairs = {}, defaultdict(list)
    for layer in sorted(os.listdir(REL_ROOT)):
        d = os.path.join(REL_ROOT, layer)
        if not os.path.isdir(d):
            continue
        for fn in sorted(os.listdir(d)):
            if not fn.endswith(".relationship.schema.json"):
                continue
            with open(os.path.join(d, fn), encoding="utf-8") as f:
                s = json.load(f)
            props = s.get("properties", {})
            rid = props.get("id", {}).get("const")
            src = props.get("source_spec_node_id", {}).get("const")
            dst = props.get("destination_spec_node_id", {}).get("const")
            pred = props.get("predicate", {}).get("const")
            if not (rid and src and dst and pred):
                continue
            desc[rid] = s.get("description", "")
            pairs[(src, dst)].append(pred)
    return desc, pairs


def suggestion(row):
    k, rep = row["kind"], row["replacement"]
    if k == "remove-duplicate":
        return "delete", f"duplicate of {rep[0]}"
    if k == "rename-predicate":
        return "rename", f"to {parts(rep[0])[1]}"
    if k == "reverse":
        return "reverse", f"to {rep[0]}" + (" (exists: delete)" if row["replacement_exists"] else "")
    if k == "retarget":
        return "retarget", f"to {rep[0]} (a migration drops the old edge)"
    if k == "compound":
        return "review", f"plan says {rep[0]}"
    return "review", "claims disagree: " + " | ".join(rep)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--layers", default="")
    ap.add_argument("--out", default=os.path.join(REFL, "walk-sheet.xlsx"))
    args = ap.parse_args()

    with open(os.path.join(REFL, "plan-clef.json"), encoding="utf-8") as f:
        plan = [r for r in json.load(f) if r.get("current_exists")]
    claims = {c["id"]: c for c in read_jsonl(CLAIMS)}
    desc, pairs = load_catalog()
    layers = [l for l in args.layers.split(",") if l] or sorted(
        {r["layer"] for r in plan if not os.path.exists(os.path.join(REFL, f"applied-walk-{r['layer']}.json"))}
    )
    kind_order = {"remove-duplicate": 0, "rename-predicate": 1, "reverse": 2, "retarget": 3, "compound": 4, "conflict": 5}

    wb = Workbook()
    how = wb.active
    how.title = "How to use"
    lines = [
        ("Relationship catalog decisions", True),
        ("One sheet per layer; one row per catalog relationship the reflection plan wants changed.", False),
        ("Fill the yellow columns only: Outcome (dropdown), Replacement (a 5-part triple, prefilled with the plan's), Also delete (more triples, comma-separated), Notes.", False),
        ("Outcome values: keep = leave the schema as is; delete = remove it (Replacement may name the edge that already covers it); rename = same pair, new predicate (Replacement names it);", False),
        ("reverse = swap source and destination (Replacement names the swapped triple); retarget = different type on one side (Replacement names it; the model migration drops the old edge); defer = decide later (row is skipped).", False),
        ("Columns to the left are read-only framing: the current relationship and its description, the other predicates the catalog has between the pair in each direction, the claims, the three sorting answers Clef gave, the plan's replacement and a mechanical suggestion.", False),
        ("Import with: python3 cli/scripts/spec-enrich/sheet-import.py --layer <id> [--sheet <file>]  ->  walk-<layer>.jsonl, then: bun apply-walk.ts --layer <id>", False),
        ("Rows already decided in walk-<layer>.jsonl appear filled in; they are the format to follow.", False),
        (f"Generated {date.today().isoformat()} from plan-clef.json; sorting answers from reflect/clef/<layer>.jsonl.", False),
    ]
    for i, (text, bold) in enumerate(lines, 1):
        c = how.cell(row=i, column=1, value=text)
        c.font = Font(name=FONT, bold=bold, size=12 if bold else 10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    how.column_dimensions["A"].width = 140

    headers = [
        ("Row", 5), ("Kind", 16), ("p", 6), ("Current relationship", 46), ("Current description", 50),
        ("Catalog: other predicates source -> destination", 22), ("Catalog: predicates destination -> source", 22),
        ("Claims", 60), ("Claims' proposed change", 44),
        ("Sorting: best source of 'pred dst'", 34), ("Sorting: best destination of 'src pred'", 34), ("Sorting: best predicate src -> dst", 34),
        ("Plan replacement", 46), ("Suggested", 10), ("Why", 36),
        ("Outcome", 11), ("Replacement", 46), ("Also delete", 30), ("Notes", 50),
    ]
    fill_in = PatternFill("solid", fgColor="FFFF00")
    head_fill = PatternFill("solid", fgColor="DDDDDD")

    for layer in layers:
        rows = sorted([r for r in plan if r["layer"] == layer], key=lambda r: (kind_order.get(r["kind"], 9), -r["p"]))
        if not rows:
            continue
        refl = defaultdict(list)
        for x in read_jsonl(os.path.join(REFL, f"{layer}.jsonl")):
            refl[x["triple"]].append(x)
        decided = {w["triple"]: w for w in read_jsonl(os.path.join(REFL, f"walk-{layer}.jsonl"))}
        ws = wb.create_sheet(layer)
        for col, (name, width) in enumerate(headers, 1):
            c = ws.cell(row=1, column=col, value=name)
            c.font = Font(name=FONT, bold=True)
            c.fill = head_fill
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.freeze_panes = "E2"
        dv = DataValidation(type="list", formula1='"' + ",".join(OUTCOMES) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        for i, r in enumerate(rows, 1):
            t = r["triple"]
            src, pred, dst = parts(t)
            others_fwd = [p for p in pairs.get((src, dst), []) if p != pred]
            back = pairs.get((dst, src), [])
            rs = refl.get(t, [])
            top = lambda q: " || ".join(x["answers"][q]["top3"] for x in rs) if rs else ""
            texts = "\n".join(f"[{cid}] {claims.get(cid, {}).get('text', '')}" for cid in r["claims"])
            proposed = "\n".join(claims.get(cid, {}).get("proposed_change", "") for cid in r["claims"])
            sug, why = suggestion(r)
            d = decided.get(t)
            values = [
                i, r["kind"], round(r["p"], 2), t, desc.get(t, ""), ", ".join(others_fwd) or "(none)", ", ".join(back) or "(none)",
                texts, proposed, top("source_for"), top("target_for"), top("predicate_for"),
                " | ".join(r["replacement"]) + (" (exists)" if r["replacement_exists"] else ""), sug, why,
                d["decision"] if d else None,
                d.get("replacement") if d else (r["replacement"][0] if len(r["replacement"]) == 1 and r["kind"] != "remove-duplicate" else None),
                ", ".join(d.get("also_delete", []) + d.get("add", [])) if d else None,
                d.get("rationale") if d else None,
            ]
            for col, v in enumerate(values, 1):
                c = ws.cell(row=i + 1, column=col, value=v)
                c.font = Font(name=FONT, size=10)
                c.alignment = Alignment(wrap_text=True, vertical="top")
                if col >= 16:
                    c.fill = fill_in
            dv.add(ws.cell(row=i + 1, column=16))
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    wb.save(args.out)
    print(f"wrote {args.out}: sheets {', '.join(layers)}; rows " + ", ".join(f"{l} {sum(1 for r in plan if r['layer'] == l)}" for l in layers))


if __name__ == "__main__":
    sys.exit(main())
